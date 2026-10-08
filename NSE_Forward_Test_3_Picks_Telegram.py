"""
NSE FORWARD TEST — 3 PICKS (OLD / NEW / V3)
--------------------------------------------

Purpose:
    This is a LIVE/FORWARD TEST tracker, NOT a historical backtest.

Every run:
    1. Uses only data available through the latest COMPLETED trading day.
    2. Selects exactly 3 stocks:
           OLD = old model
           NEW = V2 / NEW model
           V3  = V3 model
    3. Assigns the next NSE trading session as the TRADE DATE.
    4. Checks any previous forward-test rows whose next trading session
       has now occurred.
    5. Saves a permanent CSV + Excel history.
    6. Sends the Excel tracker to Telegram when Telegram credentials are configured.

The entry/T1/SL rules are kept consistent with the existing V3 backtest:
    Entry trigger = Scan Close + 0.30%
    Stop Loss     = Entry - 2.20%
    Target 1      = Entry + 2.10%

Result rules:
    NO ENTRY
    T1 HIT
    SL HIT
    AMBIGUOUS
    ENTRY, NEITHER

IMPORTANT:
    - Run after the NSE market closes for a clean completed-day scan.
    - This does not place orders.
    - It is intended to forward-test the 3 model picks objectively.
"""

import os
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
import yfinance as yf
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# SETTINGS
# ============================================================

TOP_TRACKED = 3

MIN_PRICE = 50
MAX_PRICE = 15000
MIN_VOLUME = 300_000

# Same evaluation assumptions used by your previous V3 backtest.
ENTRY_TRIGGER_PCT = 0.003
SL_PCT = 0.022
T1_PCT = 0.021

LOOKBACK_PERIOD = "2y"

CSV_FILE = "NSE_Forward_Test_3_Picks.csv"
XLSX_FILE = "NSE_Forward_Test_3_Picks.xlsx"

IST = ZoneInfo("Asia/Kolkata")


# ============================================================
# NSE UNIVERSE
# ============================================================

UNIVERSE = [
    "TCS","INFY","HCLTECH","WIPRO","TECHM","LTIM","MPHASIS","COFORGE","PERSISTENT","OFSS",
    "HDFCBANK","ICICIBANK","SBIN","KOTAKBANK","AXISBANK","BANDHANBNK","FEDERALBNK",
    "IDFCFIRSTB","INDUSINDBK","PNB","BANKBARODA","CANARABANK",
    "BAJFINANCE","BAJAJFINSV","CHOLAFIN","MUTHOOTFIN","LICHSGFIN","PFC","RECLTD",
    "BHARTIARTL","IDEA","TATACOMM","MARUTI","TATAMOTORS","M&M","BAJAJ-AUTO",
    "HEROMOTOCO","EICHERMOT","TVSMOTOR","ASHOKLEY","TATASTEEL","JSWSTEEL","HINDALCO",
    "VEDL","COALINDIA","NMDC","SAIL","NATIONALUM","SUNPHARMA","DRREDDY","CIPLA",
    "DIVISLAB","AUROPHARMA","LUPIN","BIOCON","TORNTPHARM","RELIANCE","ONGC","BPCL",
    "IOC","GAIL","POWERGRID","NTPC","ADANIGREEN","TATAPOWER","ADANIPORTS",
    "HINDUNILVR","ITC","NESTLEIND","BRITANNIA","DABUR","MARICO","GODREJCP","COLPAL",
    "ULTRACEMCO","GRASIM","AMBUJACEM","ACC","LT","SIEMENS","ABB","BHEL","TITAN",
    "ASIANPAINT","HAVELLS","VOLTAS","DMART","PIDILITIND","ADANIENT","ADANIPOWER",
    "ZOMATO","NAUKRI","IRCTC","RVNL","IRFC","RAILTEL","DELHIVERY","NYKAA","POLICYBZR",
    "PAYTM","BERGEPAINT","PREMIERPOL","MOREPENLAB","CUPID","AXISCADES","FIRSTCRY",
    "SBC","LLOYDSENGG","MARINE","SHUKRAPHAR","MANALIPETC"
]

UNIVERSE = list(dict.fromkeys(
    x.upper().replace(" ", "") for x in UNIVERSE
))

YF_ALIASES = {
    "ZOMATO": "ETERNAL",
    "CANARABANK": "CANBK",
    "TATAMOTORS": "TMCV",
}


# ============================================================
# DATA
# ============================================================

def get_history(symbol: str) -> pd.DataFrame:
    """Download daily OHLCV for one NSE symbol with retries."""
    yf_symbol = YF_ALIASES.get(symbol, symbol)

    # NIFTY index symbols already contain ^ and must not get .NS.
    ticker = yf_symbol if yf_symbol.startswith("^") else yf_symbol + ".NS"

    for attempt in range(3):
        try:
            h = yf.download(
                ticker,
                period=LOOKBACK_PERIOD,
                auto_adjust=False,
                progress=False,
                actions=False,
                threads=False,
                timeout=20,
            )

            if h is None or h.empty:
                return pd.DataFrame()

            if isinstance(h.columns, pd.MultiIndex):
                h.columns = h.columns.get_level_values(0)

            h.columns = [str(c).title() for c in h.columns]

            cols = ["Open", "High", "Low", "Close", "Volume"]
            if not all(c in h.columns for c in cols):
                return pd.DataFrame()

            h = h[cols].dropna(subset=["Close"]).copy()
            h.index = pd.to_datetime(h.index).tz_localize(None)

            return h

        except Exception as exc:
            if attempt == 2:
                print(f"  [skip] {symbol}: {exc}")
                return pd.DataFrame()
            time.sleep(1.5)

    return pd.DataFrame()


def rsi14(close: pd.Series) -> pd.Series:
    d = close.diff()
    gain = d.clip(lower=0).rolling(14).mean()
    loss = (-d.clip(upper=0)).rolling(14).mean()
    return 100 - 100 / (1 + gain / loss.replace(0, 0.001))


def atr14(h: pd.DataFrame) -> pd.Series:
    prev = h["Close"].shift(1)

    tr = pd.concat(
        [
            h["High"] - h["Low"],
            (h["High"] - prev).abs(),
            (h["Low"] - prev).abs(),
        ],
        axis=1,
    ).max(axis=1)

    return tr.rolling(14).mean()


def latest_completed_scan_date(nifty: pd.DataFrame) -> date:
    """
    Use the latest COMPLETED trading day.
    Before 4 PM IST, exclude today's possibly incomplete daily candle.
    After 4 PM, today's completed candle can be used if Yahoo already has it.
    """
    if nifty.empty:
        raise RuntimeError("Nifty data unavailable.")

    dates = sorted(set(nifty.index.date))
    today_ist = datetime.now(IST).date()
    hour_ist = datetime.now(IST).hour

    if today_ist in dates and hour_ist < 16:
        dates = [d for d in dates if d < today_ist]

    if not dates:
        raise RuntimeError("No completed trading day available.")

    return dates[-1]


# ============================================================
# MODEL 1 — OLD
# ============================================================

def old_score(sym: str, h: pd.DataFrame, asof: date):
    x = h[h.index.date <= asof]

    if len(x) < 50:
        return None

    c = float(x["Close"].iloc[-1])
    prev = float(x["Close"].iloc[-2])
    vol = float(x["Volume"].iloc[-1])

    if not (MIN_PRICE <= c <= MAX_PRICE) or vol < MIN_VOLUME:
        return None

    close = x["Close"]

    sma10 = float(close.tail(10).mean())
    sma20 = float(close.tail(20).mean())
    rsi = float(rsi14(close).iloc[-1])

    if pd.isna(rsi):
        return None

    vr = vol / max(float(x["Volume"].tail(20).mean()), 1)

    high52 = float(x["High"].tail(252).max())
    pct_high = (c - high52) / high52 * 100

    chg = (c / prev - 1) * 100

    trend = (
        "Bullish" if c > sma10 > sma20
        else "Bearish" if c < sma10 < sma20
        else "Neutral"
    )

    score = (
        3 if chg >= 2 else
        2 if chg >= 1 else
        1 if chg >= 0.5 else
        0
    )

    score += 2 if trend == "Bullish" else 0
    score += 2 if 45 < rsi < 70 else 0
    score += 2 if vr >= 1.5 else 0
    score += 1 if c > sma10 else 0
    score += 1 if pct_high > -5 else 0

    return {
        "Symbol": sym,
        "Score": round(score, 2),
        "Close": round(c, 2),
    }


# ============================================================
# MODEL 2 — NEW / V2
# ============================================================

def new_score(sym: str, h: pd.DataFrame, asof: date, nifty: pd.DataFrame):
    x = h[h.index.date <= asof]

    if len(x) < 60:
        return None

    close = x["Close"]
    vol = x["Volume"]

    c = float(close.iloc[-1])
    prev = float(close.iloc[-2])

    if not (MIN_PRICE <= c <= MAX_PRICE):
        return None

    if float(vol.iloc[-1]) < MIN_VOLUME:
        return None

    sma10 = float(close.tail(10).mean())
    sma20 = float(close.tail(20).mean())
    sma50 = float(close.tail(50).mean())

    rsi = float(rsi14(close).iloc[-1])
    atr = float(atr14(x).iloc[-1])

    if pd.isna(rsi) or pd.isna(atr) or atr <= 0:
        return None

    vr = float(vol.iloc[-1] / max(float(vol.tail(20).mean()), 1))

    mom3 = (c / float(close.iloc[-4]) - 1) * 100
    mom5 = (c / float(close.iloc[-6]) - 1) * 100

    day_range = max(
        float(x["High"].iloc[-1] - x["Low"].iloc[-1]),
        0.01,
    )

    range_pos = (
        c - float(x["Low"].iloc[-1])
    ) / day_range

    atr_pct = atr / c * 100
    dist20 = (c / sma20 - 1) * 100

    score = 0.0

    score += max(-2, min(4, mom3 * 0.9))
    score += max(-2, min(3, mom5 * 0.5))

    score += (
        2 if c > sma10 > sma20
        else 1 if c > sma20
        else 0
    )

    score += 1.5 if sma20 > sma50 else 0

    score += (
        2 if 52 <= rsi <= 68
        else 1 if 45 <= rsi < 52 or 68 < rsi <= 72
        else 0
    )

    score += (
        2 if vr >= 1.5
        else 1 if vr >= 1.1
        else 0
    )

    score += 1 if range_pos >= 0.65 else 0
    score += 1 if -3 <= dist20 <= 6 else 0

    score -= 2 if dist20 > 8 else 0

    score -= 1 if atr_pct < 1 else 0
    score -= 1 if atr_pct > 7 else 0

    rs = 0.0

    if nifty is not None and not nifty.empty and len(nifty) >= 6:
        n = nifty[nifty.index.date <= asof]

        if len(n) >= 6:
            nmom5 = (
                float(n["Close"].iloc[-1])
                / float(n["Close"].iloc[-6])
                - 1
            ) * 100

            rs = mom5 - nmom5

            score += (
                2 if rs > 2
                else 1 if rs > 0.5
                else -1 if rs < -2
                else 0
            )

    return {
        "Symbol": sym,
        "Score": round(score, 2),
        "Close": round(c, 2),
    }


# ============================================================
# MODEL 3 — V3
# ============================================================

def v3_score(sym: str, h: pd.DataFrame, asof: date, nifty: pd.DataFrame):
    x = h[h.index.date <= asof].copy()

    if len(x) < 80:
        return None

    close = x["Close"]
    vol = x["Volume"]

    c = float(close.iloc[-1])

    if not (MIN_PRICE <= c <= MAX_PRICE):
        return None

    if float(vol.iloc[-1]) < MIN_VOLUME:
        return None

    sma5 = float(close.tail(5).mean())
    sma10 = float(close.tail(10).mean())
    sma20 = float(close.tail(20).mean())
    sma50 = float(close.tail(50).mean())

    r = float(rsi14(close).iloc[-1])
    atr = float(atr14(x).iloc[-1])

    if pd.isna(r) or pd.isna(atr) or atr <= 0:
        return None

    atr_pct = atr / c * 100
    vr = float(vol.iloc[-1] / max(float(vol.tail(20).mean()), 1))

    mom1 = (c / float(close.iloc[-2]) - 1) * 100
    mom3 = (c / float(close.iloc[-4]) - 1) * 100
    mom5 = (c / float(close.iloc[-6]) - 1) * 100
    mom10 = (c / float(close.iloc[-11]) - 1) * 100

    recent = close.tail(6).diff().dropna()

    up_days = int((recent > 0).sum())
    down_days = int((recent < 0).sum())

    hi20 = float(x["High"].tail(20).max())
    lo20 = float(x["Low"].tail(20).min())

    range20 = max(hi20 - lo20, 0.01)
    range_pos20 = (c - lo20) / range20

    dist20 = (c / sma20 - 1) * 100
    near_high = (c / hi20 - 1) * 100

    score = 0.0

    score += max(-2, min(3, mom3 * 0.75))
    score += max(-2, min(3, mom5 * 0.45))
    score += max(-1.5, min(2, mom10 * 0.20))

    score += (
        2.5 if c > sma5 > sma10 > sma20
        else 1.5 if c > sma10 > sma20
        else 0.5 if c > sma20
        else 0
    )

    score += 1.5 if sma20 > sma50 else 0
    score += 1.0 if (sma10 / sma20 - 1) * 100 > 0.4 else 0

    score += (
        2.0 if 55 <= r <= 68
        else 1.0 if 50 <= r < 55 or 68 < r <= 72
        else -1.0 if r > 75
        else 0
    )

    score += (
        2.0 if vr >= 1.8
        else 1.0 if vr >= 1.25
        else 0
    )

    score += (
        1.5 if up_days >= 4
        else 0.75 if up_days == 3
        else 0
    )

    score -= 1.0 if down_days >= 4 else 0

    score += (
        1.5 if range_pos20 >= 0.80
        else 0.75 if range_pos20 >= 0.65
        else 0
    )

    score += 1.0 if near_high >= -2.0 else 0

    score += (
        1.0 if 1.5 <= atr_pct <= 4.5
        else 0.25 if 1.0 <= atr_pct < 1.5 or 4.5 < atr_pct <= 6.0
        else -1.5 if atr_pct > 7.0
        else -1.0
    )

    score += (
        1.0 if -2 <= dist20 <= 5
        else -1.5 if dist20 > 8
        else 0
    )

    rs = 0.0

    if nifty is not None and not nifty.empty and len(nifty) >= 6:
        n = nifty[nifty.index.date <= asof]

        if len(n) >= 6:
            nmom5 = (
                float(n["Close"].iloc[-1])
                / float(n["Close"].iloc[-6])
                - 1
            ) * 100

            rs = mom5 - nmom5

            score += (
                2.0 if rs > 2.0
                else 1.0 if rs > 0.75
                else -1.0 if rs < -2.0
                else 0
            )

    # Avoid chasing a single abnormal daily jump.
    if mom1 > 5:
        score -= 2.0
    elif mom1 > 3:
        score -= 1.0

    return {
        "Symbol": sym,
        "Score": round(score, 2),
        "Close": round(c, 2),
    }


# ============================================================
# FORWARD TEST EVALUATION
# ============================================================

def evaluate_next_session(h: pd.DataFrame, scan_date: date, scan_close: float):
    """
    Same next-session evaluation logic used in the existing backtest.

    Entry trigger:
        scan_close * (1 + 0.003)

    SL:
        entry * (1 - 0.022)

    T1:
        entry * (1 + 0.021)
    """
    future = h[h.index.date > scan_date]

    if future.empty:
        return {
            "Trade Date": "",
            "Entry": round(scan_close * (1 + ENTRY_TRIGGER_PCT), 2),
            "SL": "",
            "T1": "",
            "High": "",
            "Low": "",
            "Result": "PENDING",
        }

    d = future.iloc[0]
    trade_date = future.index[0].date()

    entry = scan_close * (1 + ENTRY_TRIGGER_PCT)
    sl = entry * (1 - SL_PCT)
    t1 = entry * (1 + T1_PCT)

    hi = float(d["High"])
    lo = float(d["Low"])

    if hi < entry:
        result = "NO ENTRY"
    else:
        hit_t1 = hi >= t1
        hit_sl = lo <= sl

        if hit_t1 and hit_sl:
            result = "AMBIGUOUS"
        elif hit_t1:
            result = "T1 HIT"
        elif hit_sl:
            result = "SL HIT"
        else:
            result = "ENTRY, NEITHER"

    return {
        "Trade Date": str(trade_date),
        "Entry": round(entry, 2),
        "SL": round(sl, 2),
        "T1": round(t1, 2),
        "High": round(hi, 2),
        "Low": round(lo, 2),
        "Result": result,
    }


# ============================================================
# LOAD / UPDATE FORWARD TRACKER
# ============================================================

BASE_COLUMNS = [
    "Scan Date",
    "Trade Date",
    "Scan Time",
    "Model",
    "Rank",
    "Symbol",
    "Score",
    "Scan Close",
    "Entry",
    "SL",
    "T1",
    "Trade High",
    "Trade Low",
    "Result",
]


def load_history() -> pd.DataFrame:
    p = Path(CSV_FILE)

    if not p.exists():
        return pd.DataFrame(columns=BASE_COLUMNS)

    try:
        df = pd.read_csv(p, dtype=str)
        for col in BASE_COLUMNS:
            if col not in df.columns:
                df[col] = ""
        return df[BASE_COLUMNS].copy()
    except Exception:
        return pd.DataFrame(columns=BASE_COLUMNS)


def update_previous_rows(history: pd.DataFrame, data: dict):
    """
    Fill results for previously scanned rows once their next trading
    session has become available.
    """
    if history.empty:
        return history

    history = history.copy()

    pending_mask = history["Result"].eq("PENDING")

    for idx in history.index[pending_mask]:

        symbol = str(history.at[idx, "Symbol"]).strip()

        try:
            scan_date = datetime.strptime(
                str(history.at[idx, "Scan Date"]),
                "%Y-%m-%d",
            ).date()

            scan_close = float(history.at[idx, "Scan Close"])

        except Exception:
            continue

        h = data.get(symbol)

        if h is None or h.empty:
            continue

        ev = evaluate_next_session(h, scan_date, scan_close)

        if ev["Result"] != "PENDING":
            history.at[idx, "Trade Date"] = ev["Trade Date"]
            history.at[idx, "Entry"] = ev["Entry"]
            history.at[idx, "SL"] = ev["SL"]
            history.at[idx, "T1"] = ev["T1"]
            history.at[idx, "Trade High"] = ev["High"]
            history.at[idx, "Trade Low"] = ev["Low"]
            history.at[idx, "Result"] = ev["Result"]

    return history


def append_or_replace_today_scan(
    history: pd.DataFrame,
    scan_rows: list[dict],
    scan_date: date,
    scan_time: str,
):
    """
    Replace today's scan rows if the script is run again on the same day.
    """
    history = history[
        history["Scan Date"].astype(str) != str(scan_date)
    ].copy()

    new_rows = []

    for row in scan_rows:
        ev = evaluate_next_session(
            row["_history"],
            scan_date,
            row["Close"],
        )

        new_rows.append({
            "Scan Date": str(scan_date),
            "Trade Date": ev["Trade Date"],
            "Scan Time": scan_time,
            "Model": row["Model"],
            "Rank": 1,
            "Symbol": row["Symbol"],
            "Score": row["Score"],
            "Scan Close": row["Close"],
            "Entry": ev["Entry"],
            "SL": ev["SL"],
            "T1": ev["T1"],
            "Trade High": ev["High"],
            "Trade Low": ev["Low"],
            "Result": ev["Result"],
        })

    if new_rows:
        history = pd.concat(
            [history, pd.DataFrame(new_rows)],
            ignore_index=True,
        )

    return history[BASE_COLUMNS]


# ============================================================
# SAVE EXCEL
# ============================================================

def save_excel(df: pd.DataFrame):
    wb = Workbook()

    ws = wb.active
    ws.title = "Forward Test"

    headers = list(df.columns)

    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A3A5C")

    for r, row in enumerate(df.itertuples(index=False), 2):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)

    ws.freeze_panes = "A2"

    if ws.max_row >= 1:
        ws.auto_filter.ref = ws.dimensions

    for c, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = min(
            max(len(str(h)) + 2, 12),
            22,
        )

    # --------------------------------------------------------
    # Comparison / summary sheet
    # --------------------------------------------------------

    summary = wb.create_sheet("Summary")

    summary["A1"] = "3-PICK FORWARD TEST SUMMARY"
    summary["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="0D1B2A")

    summary_headers = [
        "Model",
        "T1 Hits",
        "SL Hits",
        "Ambiguous",
        "No Entry",
        "Entry Neither",
        "Decisive",
        "Win Rate %",
    ]

    for c, h in enumerate(summary_headers, 1):
        cell = summary.cell(3, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A3A5C")

    for r, model in enumerate(["OLD", "NEW", "V3"], 4):
        x = df[df["Model"] == model]

        t1 = int((x["Result"] == "T1 HIT").sum())
        sl = int((x["Result"] == "SL HIT").sum())
        amb = int((x["Result"] == "AMBIGUOUS").sum())
        no_entry = int((x["Result"] == "NO ENTRY").sum())
        neither = int((x["Result"] == "ENTRY, NEITHER").sum())
        decisive = t1 + sl
        win = round(100 * t1 / decisive, 2) if decisive else 0.0

        values = [
            model,
            t1,
            sl,
            amb,
            no_entry,
            neither,
            decisive,
            win,
        ]

        for c, value in enumerate(values, 1):
            summary.cell(r, c, value)

    for c, h in enumerate(summary_headers, 1):
        summary.column_dimensions[get_column_letter(c)].width = max(
            14, len(h) + 2
        )

    # --------------------------------------------------------
    # Latest 3 picks
    # --------------------------------------------------------

    latest = wb.create_sheet("Latest 3 Picks")

    if not df.empty:
        latest_scan = df["Scan Date"].max()
        latest_df = df[df["Scan Date"] == latest_scan].copy()

        latest["A1"] = f"Latest 3 Picks — Scan Date {latest_scan}"
        latest["A1"].font = Font(bold=True, size=14, color="FFFFFF")
        latest["A1"].fill = PatternFill("solid", fgColor="0D1B2A")

        latest_headers = [
            "Model",
            "Symbol",
            "Score",
            "Scan Close",
            "Trade Date",
            "Entry",
            "SL",
            "T1",
            "Result",
        ]

        for c, h in enumerate(latest_headers, 1):
            cell = latest.cell(3, c, h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1A3A5C")

        latest_df = latest_df.sort_values(
            ["Model"],
            key=lambda s: s.map({"OLD": 1, "NEW": 2, "V3": 3}).fillna(99),
        )

        for r, row in enumerate(latest_df.to_dict("records"), 4):
            vals = [
                row["Model"],
                row["Symbol"],
                row["Score"],
                row["Scan Close"],
                row["Trade Date"],
                row["Entry"],
                row["SL"],
                row["T1"],
                row["Result"],
            ]

            for c, value in enumerate(vals, 1):
                latest.cell(r, c, value)

        for c, h in enumerate(latest_headers, 1):
            latest.column_dimensions[get_column_letter(c)].width = max(
                14, len(h) + 2
            )

    wb.save(XLSX_FILE)



# ============================================================
# TELEGRAM
# ============================================================

def send_telegram_document(file_path: str, caption: str = "") -> bool:
    """
    Send the generated Excel file to Telegram.

    Set these Render environment variables:
        TELEGRAM_BOT_TOKEN
        TELEGRAM_CHAT_ID
    """
    token = os.getenv("TELEGRAM_BOT_TOKEN")
    chat_id = os.getenv("TELEGRAM_CHAT_ID")

    if not token or not chat_id:
        print(
            "  [Telegram] TELEGRAM_BOT_TOKEN or TELEGRAM_CHAT_ID "
            "is not set; skipping upload."
        )
        return False

    url = f"https://api.telegram.org/bot{token}/sendDocument"

    try:
        with open(file_path, "rb") as fh:
            response = requests.post(
                url,
                data={
                    "chat_id": chat_id,
                    "caption": caption[:1024],
                },
                files={
                    "document": (
                        Path(file_path).name,
                        fh,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
                timeout=60,
            )

        response.raise_for_status()
        payload = response.json()

        if payload.get("ok"):
            print("  [Telegram] Excel file sent successfully.")
            return True

        print(f"  [Telegram] API returned an error: {payload}")
        return False

    except Exception as exc:
        print(f"  [Telegram] Upload failed: {exc}")
        return False


# ============================================================
# MAIN
# ============================================================

def main():
    print("\n" + "=" * 80)
    print("NSE FORWARD TEST — OLD + NEW/V2 + V3")
    print("=" * 80)

    print("\nDownloading daily data...")

    data = {}

    for i, symbol in enumerate(UNIVERSE, 1):

        h = get_history(symbol)

        if not h.empty:
            data[symbol] = h

        if i % 10 == 0 or i == len(UNIVERSE):
            print(f"  {i:>3}/{len(UNIVERSE)} symbols processed")

        time.sleep(0.05)

    nifty = get_history("^NSEI")

    if nifty.empty:
        raise RuntimeError(
            "NIFTY 50 data could not be downloaded. "
            "Please retry later."
        )

    scan_date = latest_completed_scan_date(nifty)
    scan_time = datetime.now(IST).strftime("%Y-%m-%d %H:%M:%S IST")

    print(f"\nSCAN DATE  : {scan_date}")
    print("TRADE DATE : NEXT NSE TRADING SESSION")

    # --------------------------------------------------------
    # First update all previous PENDING rows.
    # --------------------------------------------------------

    history = load_history()
    history = update_previous_rows(history, data)

    # --------------------------------------------------------
    # Calculate today's 3 picks.
    # --------------------------------------------------------

    old_rows = []
    new_rows = []
    v3_rows = []

    n = nifty[nifty.index.date <= scan_date]

    for symbol, h in data.items():

        a = old_score(symbol, h, scan_date)
        b = new_score(symbol, h, scan_date, n)
        c = v3_score(symbol, h, scan_date, n)

        if a:
            a["_history"] = h
            old_rows.append(a)

        if b:
            b["_history"] = h
            new_rows.append(b)

        if c:
            c["_history"] = h
            v3_rows.append(c)

    if not old_rows:
        raise RuntimeError("No OLD candidates passed the filters.")

    if not new_rows:
        raise RuntimeError("No NEW/V2 candidates passed the filters.")

    if not v3_rows:
        raise RuntimeError("No V3 candidates passed the filters.")

    old_pick = max(
        old_rows,
        key=lambda z: (z["Score"], z["Close"]),
    )

    new_pick = max(
        new_rows,
        key=lambda z: (z["Score"], z["Close"]),
    )

    v3_pick = max(
        v3_rows,
        key=lambda z: (z["Score"], z["Close"]),
    )

    # --------------------------------------------------------
    # Build exactly 3 forward-test rows.
    # --------------------------------------------------------

    scan_rows = [
        {
            "Model": "OLD",
            "Symbol": old_pick["Symbol"],
            "Score": old_pick["Score"],
            "Close": old_pick["Close"],
            "_history": old_pick["_history"],
        },
        {
            "Model": "NEW",
            "Symbol": new_pick["Symbol"],
            "Score": new_pick["Score"],
            "Close": new_pick["Close"],
            "_history": new_pick["_history"],
        },
        {
            "Model": "V3",
            "Symbol": v3_pick["Symbol"],
            "Score": v3_pick["Score"],
            "Close": v3_pick["Close"],
            "_history": v3_pick["_history"],
        },
    ]

    history = append_or_replace_today_scan(
        history,
        scan_rows,
        scan_date,
        scan_time,
    )

    history = history.sort_values(
        ["Scan Date", "Model"],
        ascending=[True, True],
    ).reset_index(drop=True)

    # --------------------------------------------------------
    # Save.
    # --------------------------------------------------------

    history.to_csv(CSV_FILE, index=False)
    save_excel(history)

    # Send the Excel tracker to Telegram.
    v2_row = next(r for r in scan_rows if r["Model"] == "NEW")
    v3_row = next(r for r in scan_rows if r["Model"] == "V3")
    old_row = next(r for r in scan_rows if r["Model"] == "OLD")

    telegram_caption = (
        f"NSE 3-Pick Forward Test\n"
        f"Scan Date: {scan_date}\n"
        f"OLD: {old_row['Symbol']}\n"
        f"NEW/V2: {v2_row['Symbol']}\n"
        f"V3: {v3_row['Symbol']}\n"
        f"Next Trade Session: {evaluate_next_session(v2_row['_history'], scan_date, v2_row['Close'])['Trade Date']}"
    )

    send_telegram_document(XLSX_FILE, telegram_caption)

    # --------------------------------------------------------
    # Console output.
    # --------------------------------------------------------

    print("\n" + "=" * 80)
    print("TODAY'S 3 FORWARD-TEST PICKS")
    print("=" * 80)

    for row in scan_rows:
        ev = evaluate_next_session(
            row["_history"],
            scan_date,
            row["Close"],
        )

        print(
            f"{row['Model']:>4} | "
            f"{row['Symbol']:<15} | "
            f"Score {row['Score']:>6} | "
            f"Scan Close {row['Close']:>10.2f} | "
            f"Entry {ev['Entry']:>10.2f} | "
            f"Trade Date {ev['Trade Date']}"
        )

    print("\n" + "-" * 80)
    print("Current forward-test result counts")
    print("-" * 80)

    for model in ["OLD", "NEW", "V3"]:
        x = history[history["Model"] == model]

        t1 = int((x["Result"] == "T1 HIT").sum())
        sl = int((x["Result"] == "SL HIT").sum())
        amb = int((x["Result"] == "AMBIGUOUS").sum())
        no_entry = int((x["Result"] == "NO ENTRY").sum())
        neither = int((x["Result"] == "ENTRY, NEITHER").sum())
        pending = int((x["Result"] == "PENDING").sum())

        decisive = t1 + sl
        win = round(100 * t1 / decisive, 2) if decisive else 0.0

        print(
            f"{model}: "
            f"T1={t1} | SL={sl} | "
            f"AMB={amb} | NO ENTRY={no_entry} | "
            f"NEITHER={neither} | PENDING={pending} | "
            f"Win={win}%"
        )

    print("\nFiles:")
    print(f"  CSV : {Path(CSV_FILE).resolve()}")
    print(f"  XLSX: {Path(XLSX_FILE).resolve()}")

    print("\nForward-test complete.")


if __name__ == "__main__":
    main()
